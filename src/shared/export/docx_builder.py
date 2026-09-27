import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT


def build_resume_docx(data: dict) -> io.BytesIO:
    """
    Builds a high-density, ATS-compatible resume in .docx format.
    Uses right-aligned tab stops to keep dates and locations flush right.
    """
    doc = Document()

    # 0.5 - 0.6 inch margins for optimal 1-2 page density
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Base typography: Times New Roman
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    # Candidate Name & Contact Block
    if data.get("full_name"):
        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_name.paragraph_format.space_before = Pt(0)
        p_name.paragraph_format.space_after = Pt(2)
        run_name = p_name.add_run(data["full_name"].upper())
        run_name.bold = True
        run_name.font.size = Pt(16)

    if data.get("contact_info"):
        p_contact = doc.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_contact.paragraph_format.space_before = Pt(0)
        p_contact.paragraph_format.space_after = Pt(8)
        run_contact = p_contact.add_run(data["contact_info"])
        run_contact.font.size = Pt(9.5)

    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)

    # Summary
    if data.get("summary"):
        add_section_header("Professional Summary")
        p = doc.add_paragraph(data["summary"])
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

    # Experience
    if data.get("experience"):
        add_section_header("Professional Experience")
        for exp in data["experience"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1)

            # Left flush: Title — Company
            title_run = p.add_run(f"{exp.get('title', '')} — {exp.get('company', '')}")
            title_run.bold = True

            # Right flush: Dates | Location via Tab Stop
            location_str = f" | {exp['location']}" if exp.get("location") else ""
            meta_str = f"\t{exp.get('dates', '')}{location_str}"
            meta_run = p.add_run(meta_str)
            meta_run.italic = True

            # 7.3 inches maps to the exact right margin border
            p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)

            for bullet in exp.get("bullets", []):
                bp = doc.add_paragraph(bullet, style="List Bullet")
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(1.5)
                bp.paragraph_format.line_spacing = 1.1

    # Skills
    if data.get("skills"):
        add_section_header("Technical Skills")
        for skill in data["skills"]:
            p = doc.add_paragraph(f"•  {skill}")
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1.5)

    # Education
    if data.get("education"):
        add_section_header("Education")
        for edu in data["education"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            deg = p.add_run(f"{edu.get('degree', '')} — {edu.get('institution', '')}")
            deg.bold = True
            if edu.get("dates"):
                d_run = p.add_run(f" ({edu['dates']})")
                d_run.italic = True

    # Certifications
    if data.get("certifications"):
        add_section_header("Certifications")
        p = doc.add_paragraph(f"•  {', '.join(data['certifications'])}")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_cover_letter_docx(text: str) -> io.BytesIO:
    """
    Builds a clean, modern cover letter in .docx format tailored for tech roles.
    Uses Calibri with standard 1-inch margins.
    """
    doc = Document()

    # Standard 1-inch letter margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    # Split text into paragraphs and render with balanced spacing
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    for p_text in paragraphs:
        p = doc.add_paragraph(p_text)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.2

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer